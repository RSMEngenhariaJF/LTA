"""Gera doc/Manual_LTA_Linhas_Periodicas.docx (aula + manual do software).

Uso (na raiz do projeto):
    python doc/fontes/capturar_telas.py     # opcional: atualiza as capturas de tela
    python doc/fontes/gerar_manual.py       # gera figuras, números e o .docx
Requer: python-docx, latex2mathml, Microsoft Office (conversor de equações) e,
para atualizar o sumário e gerar o PDF, pywin32.
"""
import sys
from datetime import date
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
sys.path.insert(0, str(AQUI))

import figuras  # noqa: E402
from docx_util import AZUL, CINZA, Manual  # noqa: E402

IMG = AQUI / "img"
SAIDA = RAIZ / "doc" / "Manual_LTA_Linhas_Periodicas.docx"
LIVRO = "Martín (2015)"


def n(x, d=2):
    """Número no formato brasileiro."""
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".").replace("-", "−")


def ref(secao, paginas):
    return f"{LIVRO}, seção {secao}, pp. {paginas}"


def capa(m):
    d = m.doc
    for _ in range(5):
        d.add_paragraph()
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Linhas de Transmissão Artificiais Periódicas")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), AZUL
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Teoria, projeto e uso da ferramenta LTA")
    r.font.size, r.font.color.rgb = Pt(15), CINZA
    d.add_paragraph()
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Material didático em formato de aula, baseado no capítulo 2 de F. Martín, "
                  "Artificial Transmission Lines for RF and Microwave Applications (Wiley, 2015)")
    r.italic, r.font.size = True, Pt(11)
    for _ in range(8):
        d.add_paragraph()
    info = [("Instituição", "Instituto Hardware BR (HBR)"),
            ("Classificação", "Uso interno – P&D"),
            ("Versão", f"1.0 – rascunho para revisão ({date.today().strftime('%m/%Y')})")]
    for k, v in info:
        par = d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        r = par.add_run(f"{k}: ")
        r.bold = True
        par.add_run(v)
    d.add_paragraph()
    m.tabela(["Etapa", "Função / cargo", "Data", "Assinatura"],
             [["Elaboração", "", "", ""], ["Revisão técnica", "", "", ""], ["Aprovação", "", "", ""]],
             [3.2, 5.6, 2.6, 4.6])
    m.nota("este material foi redigido com apoio de IA e ainda não passou por revisão técnica. "
           "Conforme o SGQ, deve ser revisado e aprovado por pessoa qualificada diferente de quem o "
           "elaborou antes de ser usado como referência de projeto. Os resultados numéricos vêm de "
           "modelos de circuito e devem ser confirmados por simulação eletromagnética e medição.",
           "Situação do documento")


# =============================================================================
def cap1(m):
    m.h("1  Introdução: por que linhas artificiais?")
    m.p("Linhas de transmissão são o elemento mais básico da engenharia de RF e micro-ondas. Além de "
        "transportar energia e informação entre dois pontos, elas são os blocos com que se constroem "
        "filtros, divisores, acopladores, redes de casamento e ressonadores. Em tecnologia planar (microstrip, "
        "coplanar), basta gravar trilhas numa placa de circuito impresso para obter funções que, em baixa "
        "frequência, exigiriam indutores e capacitores discretos. Esses circuitos são chamados de "
        "**distribuídos**, porque cada trecho de linha é descrito por parâmetros distribuídos ao longo do "
        "comprimento, e não concentrados em componentes.")
    m.p("Os circuitos distribuídos convencionais, porém, têm duas limitações de nascença, que o autor do "
        "livro de referência destaca logo no prefácio (" + LIVRO + ", pp. xiii–xiv):")
    m.lista(["**O tamanho acompanha o comprimento de onda.** Trechos típicos têm λ/4 ou λ/2. A 1 GHz, "
             "num substrato comum, λ/4 passa de 4 cm, o que é grande para muitos produtos.",
             "**A linha convencional oferece poucos graus de liberdade.** Só controlamos a impedância "
             "característica e o comprimento elétrico, e a fase cresce linearmente com a frequência. Um "
             "circuito projetado para uma frequência perde a fase correta quando a frequência muda, o que "
             "limita a banda e impede, por exemplo, operar em duas bandas não harmônicas."])
    m.p("**Linhas de transmissão artificiais** são a resposta a essas limitações. A ideia é modificar a "
        "linha de forma controlada, carregando-a periodicamente com elementos reativos (capacitores, "
        "indutores, ressonadores) ou perturbando periodicamente sua geometria (largura da trilha, furos no "
        "plano de terra). A estrutura resultante continua se comportando como uma linha de transmissão, mas "
        "agora com **impedância e dispersão projetáveis**: podemos torná-la mais lenta (e portanto mais "
        "curta), criar faixas de frequência em que nada se propaga, ou obter comportamentos que nenhuma "
        "linha comum apresenta.")
    m.h("1.1  Objetivos desta aula", 2)
    m.p("Ao final deste material, espera-se que o leitor consiga:")
    m.lista(["entender o que muda numa linha quando ela se torna periódica (bandas passantes e "
             "proibidas, velocidade de fase reduzida, impedância de Bloch);",
             "calcular a dispersão e a impedância de uma célula unitária pelo método da matriz ABCD;",
             "projetar uma linha com carga capacitiva (*slow-wave*) para miniaturização e supressão de "
             "espúrios;",
             "projetar uma linha EBG (*electromagnetic bandgap*) por modos acoplados para rejeitar uma "
             "faixa de frequência;",
             "usar a ferramenta **LTA** para fazer esses cálculos, visualizar a estrutura e exportá-la para "
             "KiCad, Altium, ADS, HFSS e CST."])
    m.h("1.2  Como este material está organizado", 2)
    m.p("O capítulo 2 revisa, de forma breve, os parâmetros de uma linha de transmissão e da microstrip. "
        "O capítulo 3 apresenta a ideia de estrutura periódica (teorema de Floquet e condição de Bragg). Os "
        "capítulos 4, 5 e 6 tratam, cada um, de uma ferramenta de análise e de um tipo de estrutura: a "
        "célula ABCD genérica, a linha com carga capacitiva e a linha EBG analisada por modos acoplados. "
        "Cada um desses capítulos termina com uma seção sobre **o que a estrutura modifica e quais os "
        "impactos de cada parâmetro**, com aplicações descritas na literatura, e mostra a aba "
        "correspondente do software. O capítulo 7 descreve o fluxo de trabalho e a exportação. O capítulo 8 "
        "traz exemplos de aplicação resolvidos passo a passo e o capítulo 9 discute limitações e próximos "
        "passos.")
    m.nota("a teoria é apresentada de forma resumida, apenas para dar o fio da meada. Para cada "
           "assunto indicamos a seção e as páginas do livro de F. Martín em que o conteúdo é desenvolvido "
           "com mais profundidade. As equações citadas como \"Martín (2.21)\" seguem a numeração do livro.",
           "Como ler")


# =============================================================================
def cap2(m, F):
    m.h("2  Parâmetros fundamentais de uma linha de transmissão")
    m.p("Antes de tornar uma linha periódica, vale relembrar como se descreve uma linha comum. Um trecho "
        "infinitesimal de comprimento dz é representado por uma indutância série L·dz, uma resistência série "
        "R·dz, uma capacitância em derivação C·dz e uma condutância G·dz, todas por unidade de comprimento. "
        "Desse modelo saem as duas grandezas que caracterizam qualquer linha: a constante de propagação e a "
        "impedância característica (" + ref("1.2", "5–8") + ").")
    m.eq(r"\gamma=\alpha+j\beta=\sqrt{(R+j\omega L)(G+j\omega C)}", "2.1")
    m.eq(r"Z_{0}=\sqrt{\frac{R+j\omega L}{G+j\omega C}}", "2.2")
    m.p("A parte real α é a **constante de atenuação** (em néper por metro) e a parte imaginária β é a "
        "**constante de fase** (em radiano por metro). Numa linha sem perdas (R = G = 0), as expressões se "
        "simplificam para:")
    m.eq(r"\beta=\omega\sqrt{LC},\qquad Z_{0}=\sqrt{\frac{L}{C}},\qquad v_{p}=\frac{\omega}{\beta}=\frac{1}{\sqrt{LC}}", "2.3")
    m.p("Três observações guiarão todo o restante da aula:")
    m.lista(["Na linha comum, β é **proporcional** à frequência: o gráfico de β em função de ω é uma reta. "
             "Dizemos que a linha não é dispersiva. Toda a graça das linhas artificiais está em entortar "
             "essa reta.",
             "Z₀ depende apenas da relação L/C, que é fixada pela geometria e pelo substrato. Não depende da "
             "frequência.",
             "A velocidade de fase v_p = 1/√(LC) é o que define o comprimento de onda guiado λ_g = v_p/f e, "
             "portanto, o tamanho físico dos circuitos. Aumentar L ou C por unidade de comprimento deixa a "
             "onda mais lenta e o circuito mais curto: esse é o princípio da miniaturização por "
             "*slow-wave*."])
    m.h("2.1  A linha microstrip", 2)
    m.p("Na microstrip, uma trilha de largura W fica sobre um dielétrico de espessura h e permissividade "
        "εr, com plano de terra embaixo. Como parte do campo fica no ar e parte no dielétrico, a onda "
        "\"enxerga\" uma **permissividade efetiva** εeff, intermediária entre 1 e εr. A velocidade e o "
        "comprimento de onda guiado ficam:")
    m.eq(r"v_{p}=\frac{c}{\sqrt{\varepsilon_{\mathrm{eff}}}},\qquad\lambda_{g}=\frac{\lambda_{0}}{\sqrt{\varepsilon_{\mathrm{eff}}}}", "2.4")
    m.p("O software LTA calcula εeff e Z₀ pelas fórmulas fechadas de Hammerstad e Jensen (1980), "
        "reproduzidas a seguir com u = W/h e η₀ ≈ 376,73 Ω:")
    m.eq(r"\varepsilon_{\mathrm{eff}}=\frac{\varepsilon_{r}+1}{2}+\frac{\varepsilon_{r}-1}{2}\left(1+\frac{10}{u}\right)^{-ab}", "2.5")
    m.eq(r"a=1+\frac{1}{49}\ln\left(\frac{u^{4}+(u/52)^{2}}{u^{4}+0,432}\right)+\frac{1}{18,7}\ln\left(1+\left(\frac{u}{18,1}\right)^{3}\right),\quad b=0,564\left(\frac{\varepsilon_{r}-0,9}{\varepsilon_{r}+3}\right)^{0,053}", "2.6")
    m.eq(r"Z_{0}=\frac{\eta_{0}}{2\pi\sqrt{\varepsilon_{\mathrm{eff}}}}\ln\left(\frac{f(u)}{u}+\sqrt{1+\frac{4}{u^{2}}}\right),\quad f(u)=6+(2\pi-6)e^{-(30,666/u)^{0,7528}}", "2.7")
    m.p("Para obter a largura W a partir de uma impedância desejada (síntese), o software inverte "
        "numericamente a Eq. (2.7) por bisseção. Essas fórmulas desprezam a espessura do cobre e a "
        "dispersão da microstrip, aproximação de primeira ordem também adotada no livro. Como referência, "
        f"uma linha de 50 Ω em RO3010 (εr = 10,2, h = 1,27 mm) resulta em W = {n(F['livro_ebg']['w50'], 3)} mm, "
        "contra 1,2 mm citado no livro (Fig. 2.16).")
    m.figura(IMG / "tela_01_substrato.png", "2.1",
             "Aba Substrato do LTA: os parâmetros definidos aqui valem para todas as abas. A calculadora "
             "converte Z₀ ↔ W e o gráfico mostra Z₀ e εeff em função de W/h.")
    m.p("A comparação entre microstrip, coplanar e outras tecnologias planares, bem como o tratamento de "
        "perdas (tangente de perdas e profundidade pelicular), estão em " + ref("1.4 e 1.5", "16–31") + ".")
    m.h("2.2  Parâmetros S e matriz ABCD", 2)
    m.p("Duas descrições de quadripolos serão usadas o tempo todo. A **matriz de espalhamento** (parâmetros "
        "S) relaciona ondas incidentes e refletidas nas portas: |S₁₁| indica quanto é refletido e |S₂₁| quanto "
        "é transmitido. É o que se mede num analisador de redes e o que se exporta em arquivos Touchstone. A "
        "**matriz ABCD** (ou de transmissão) relaciona tensão e corrente na entrada com as da saída:")
    m.eq(r"\begin{bmatrix}V_{1}\\I_{1}\end{bmatrix}=\begin{bmatrix}A&B\\C&D\end{bmatrix}\begin{bmatrix}V_{2}\\I_{2}\end{bmatrix}", "2.8")
    m.p("A grande vantagem da ABCD é que a matriz de uma cascata é o **produto** das matrizes de cada "
        "trecho. Por isso ela é a ferramenta natural para estruturas periódicas. Para uma rede recíproca, "
        "AD − BC = 1. A conversão para parâmetros S com impedância de referência Z₀ nas duas portas é "
        "(Pozar, 2012, tab. 4.2):")
    m.eq(r"S_{11}=\frac{A+B/Z_{0}-CZ_{0}-D}{\Delta},\qquad S_{21}=\frac{2}{\Delta},\qquad\Delta=A+\frac{B}{Z_{0}}+CZ_{0}+D", "2.9")
    m.p("A matriz de espalhamento e suas propriedades estão resumidas no Apêndice C do livro "
        "(" + LIVRO + ", pp. 474–479).")


# =============================================================================
def cap3(m, F):
    m.h("3  Estruturas periódicas: Floquet, Bragg e bandas proibidas")
    m.p("Considere agora uma linha que se repete a cada período l: ou porque recebeu um elemento "
        "idêntico a cada l, ou porque sua geometria varia periodicamente. O **teorema de Floquet** "
        "afirma que, numa estrutura periódica infinita, os campos (ou a tensão e a corrente) podem ser "
        "escritos como uma onda que se propaga multiplicada por uma função que se repete a cada período "
        "(" + ref("2.2", "48–53") + "):")
    m.eq(r"\psi(z)=e^{-\gamma z}P(z),\qquad P(z+l)=P(z)", "3.1")
    m.p("A consequência prática é simples e poderosa: de uma célula para a seguinte, a onda só muda por um "
        "fator constante, e^(−γl). Esse tipo de onda é chamado de **onda de Bloch**. Toda a física da linha "
        "periódica fica contida no valor de γ, que agora é uma função complicada da frequência.")
    m.h("3.1  Harmônicos espaciais e a condição de Bragg", 2)
    m.p("Como P(z) é periódica, pode ser expandida em série de Fourier. O campo passa a ser uma soma de "
        "ondas, os **harmônicos espaciais**, com constantes de fase deslocadas de múltiplos de 2π/l:")
    m.eq(r"\beta_{n}=\beta+\frac{2\pi n}{l},\qquad n=0,\pm1,\pm2,\ldots", "3.2")
    m.p("Os harmônicos não existem isoladamente: são partes de um mesmo modo. Mas eles explicam o "
        "fenômeno central. Onde a reta β(ω) de um harmônico cruza a de outro que viaja em sentido "
        "contrário, os dois se acoplam fortemente, a energia da onda direta é transferida para a onda "
        "refletida e surge uma **banda proibida** (*stop band*, *band gap*). O primeiro cruzamento ocorre "
        "quando βl = π, isto é, quando o período vale meio comprimento de onda guiado. Essa é a **condição "
        "de Bragg**:")
    m.eq(r"f_{\mathrm{Bragg}}=\frac{v_{p}}{2l}=\frac{c}{2l\sqrt{\varepsilon_{\mathrm{eff}}}}", "3.3")
    m.p("e se repete nos múltiplos inteiros dessa frequência, desde que a perturbação tenha conteúdo "
        "harmônico correspondente (voltaremos a isso no capítulo 6). A Figura 3.1 ilustra o efeito para uma "
        "linha levemente carregada: a curva da linha periódica (azul) acompanha a da linha comum (pontilhada, "
        "rebatida para dentro do intervalo 0 ≤ βl ≤ π) e se separa dela justamente nos pontos de Bragg, "
        "abrindo faixas em que a atenuação por célula, αl, deixa de ser nula.")
    m.figura(IMG / "fig_dispersao_conceito.png", "3.1",
             f"Linha com carga periódica leve (calculada com o LTA). (a) Diagrama de dispersão na zona "
             f"reduzida; as linhas tracejadas marcam múltiplos de f = {n(F['conceito']['fs'] / 1e9, 1)} GHz, "
             "onde βl = nπ. (b) Atenuação por célula: as bandas proibidas aparecem nos pontos de Bragg.")
    m.h("3.2  Como ler um diagrama de dispersão", 2)
    m.p("O diagrama β × f (ou de Brillouin) é a \"impressão digital\" de uma linha periódica. Ao longo da "
        "aula usaremos as seguintes leituras:")
    m.lista(["**Zona de Brillouin:** o diagrama se repete a cada 2π em βl, então basta desenhá-lo em "
             "0 ≤ βl ≤ π (zona reduzida).",
             "**Inclinação da curva:** é a velocidade de grupo. Na origem, a razão ω/β dá a velocidade de "
             "fase em baixa frequência. Curva mais deitada (mais βl para a mesma frequência) significa "
             "**onda mais lenta**, ou seja, linha eletricamente mais longa por milímetro.",
             "**Bandas proibidas:** faixas de frequência em que não há solução com α = 0. Ali a onda é "
             "evanescente e a estrutura reflete a potência de volta.",
             "**Ondas rápidas (v_p > c):** em estruturas abertas, permitem radiação (antenas *leaky-wave*), "
             "tema tratado nas pp. 51–53 e no capítulo 4 do livro."])
    m.nota("a relação entre γ, a impedância e os parâmetros S de uma estrutura finita é o que "
           "faremos nos próximos capítulos. Primeiro com elementos concentrados (matriz ABCD, cap. 4 e 5), "
           "depois com perturbações contínuas da geometria (modos acoplados, cap. 6).", "Para onde vamos")


# =============================================================================
def cap4(m, F):
    m.h("4  O método da matriz de transferência (ABCD)")
    m.p("Quando a célula unitária pode ser descrita por uma matriz ABCD (o que vale para qualquer "
        "combinação de elementos concentrados e trechos de linha), a aplicação do teorema de Floquet é "
        "imediata (" + ref("2.3.1", "54–55") + "). As tensões e correntes nos terminais de duas células "
        "vizinhas obedecem, ao mesmo tempo, à matriz da célula e ao fator de propagação de Bloch:")
    m.eq(r"\begin{bmatrix}V_{n}\\I_{n}\end{bmatrix}=\begin{bmatrix}A&B\\C&D\end{bmatrix}\begin{bmatrix}V_{n+1}\\I_{n+1}\end{bmatrix},\qquad V_{n+1}=e^{-\gamma l}V_{n},\quad I_{n+1}=e^{-\gamma l}I_{n}", "4.1")
    m.p("Trata-se de um problema de autovalores: os autovetores são os pares (V, I) que se propagam e os "
        "autovalores dão o fator de propagação. Impondo determinante nulo e usando a reciprocidade "
        "(AD − BC = 1), chega-se à **relação de dispersão** (Martín 2.21):")
    m.eq(r"\cosh(\gamma l)=\frac{A+D}{2}", "4.2")
    m.p("Num sistema sem perdas, o lado direito é real. Há então duas possibilidades: se |A + D|/2 ≤ 1, "
        "γ = jβ é imaginário puro e a onda se propaga (banda passante), com cos(βl) = (A + D)/2; se "
        "|A + D|/2 > 1, γ = α é real e a onda é atenuada (banda proibida). Para células simétricas "
        "(A = D), a relação fica simplesmente cos(βl) = A.")
    m.h("4.1  Impedância de Bloch", 2)
    m.p("A razão entre tensão e corrente de uma onda de Bloch faz o papel da impedância característica "
        "da linha periódica e é chamada de **impedância de Bloch** (" + ref("2.3.2", "56–59") + "):")
    m.eq(r"Z_{B}^{\pm}=\frac{-2B}{A-D\mp\sqrt{(A+D)^{2}-4}}", "4.3")
    m.p("O sinal + corresponde à onda que viaja no sentido positivo e o sinal − à onda contrária. Para "
        "célula simétrica, a expressão se reduz a Z_B = ±B/√(A² − 1). Diferentemente de Z₀ numa linha "
        "comum, Z_B **depende da frequência**: é real na banda passante e imaginária na banda proibida "
        "(célula simétrica sem perdas). Uma consequência importante, demonstrada no livro, é que a impedância "
        "de Bloch coincide com a **impedância iterativa** (aquela que, ligada na saída, é vista na "
        "entrada) para qualquer número de células, mesmo uma só (" + ref("2.3.5", "63–64") + ").")
    m.p("Na prática, isso quer dizer que, **em cada frequência da banda passante, a linha periódica se "
        "comporta como uma linha comum** com constante de fase β e impedância Z_B. Se a carga for igual a "
        "Z_B naquela frequência, toda a potência é entregue; se o comprimento elétrico for múltiplo de π, a "
        "entrada \"vê\" a própria carga (" + ref("2.3.4", "62–63") + ").")
    m.h("4.2  Exemplo: a linha comum como cascata de células LC", 2)
    m.p("O exemplo clássico, analisado no livro (" + ref("2.3.3", "60–62") + "), é aproximar uma linha sem "
        "perdas por uma cascata de células com indutância série L e capacitância em derivação C. Para a "
        "célula simétrica em T (L/2 – C – L/2) com Z_s = jωL e Z_p = 1/(jωC), a Eq. (4.2) e a Eq. (4.3) dão:")
    m.eq(r"\cos(\beta l)=1-\frac{\omega^{2}LC}{2},\qquad Z_{B}=\sqrt{\frac{L}{C}}\sqrt{1-\left(\frac{\omega}{\omega_{0}}\right)^{2}},\qquad\omega_{0}=\frac{2}{\sqrt{LC}}", "4.4")
    m.p("Com L = 2,5 nH e C = 1 pF, Z_B em baixa frequência vale √(L/C) = 50 Ω e a frequência de corte é "
        "f₀ = 1/(π√(LC)) ≈ 6,37 GHz. Abaixo dela, a cascata se comporta como uma linha de 50 Ω; acima, "
        "torna-se um filtro passa-baixas. A lição é dupla: (i) a cascata de células só imita a linha "
        "enquanto cada célula for eletricamente curta; (ii) o **corte** é uma propriedade intrínseca da "
        "periodicidade, que será explorada no capítulo 5.")
    m.p("E se a mesma rede for descrita por células em L (L – C), assimétricas? A dispersão é idêntica, "
        "pois só depende de A + D, mas a impedância de Bloch ganha uma parte imaginária ±jωL/2 (Martín "
        "2.32). Em uma estrutura finita, terminada em 50 Ω, as respostas das duas descrições divergem à "
        "medida que a frequência sobe. Esse é um alerta prático: ao modelar uma estrutura real por células, "
        "a escolha dos planos de referência precisa respeitar a assimetria física da célula.")
    m.h("4.3  A aba Célula ABCD do LTA", 2)
    m.p("A aba **Célula ABCD** implementa exatamente esse raciocínio para uma célula genérica. O ramo série "
        "é uma indutância Ls em série com uma capacitância Cs; o ramo em derivação é uma capacitância Cp em "
        "paralelo com uma indutância Lp. Elemento com valor zero é tratado como ausente (Cs = 0 vira curto "
        "e Lp = 0 vira circuito aberto). Com isso a mesma aba cobre a linha convencional (Ls, Cp), a linha "
        "*left-handed* (Cs, Lp) e a linha CRLH completa (os quatro elementos), que é o assunto do capítulo 3 "
        "do livro. As topologias disponíveis são T, Π, L e L invertida, e é possível comparar duas delas "
        "lado a lado.")
    m.figura(IMG / "tela_02_celula_estrutura.png", "4.1",
             "Aba Célula ABCD, sub-aba Estrutura: o esquemático da célula é redesenhado enquanto os valores "
             "são editados, mostrando o valor de cada elemento (por exemplo, Ls/2 em cada braço da célula T).")
    m.figura(IMG / "tela_03_celula_resultados.png", "4.2",
             "Resultados do exemplo do livro (Fig. 2.8): L = 2,5 nH, C = 1 pF, 20 células, células T e L "
             "comparadas. Dispersão idêntica, impedâncias de Bloch com partes imaginárias diferentes e "
             "respostas S que divergem com a frequência.")
    m.p("Os quatro gráficos de resultado são: o diagrama de dispersão (βl/π e αl/π no eixo horizontal, "
        "frequência no vertical, como no livro); as partes real e imaginária de Z_B; e |S₂₁| e |S₁₁| da "
        "cascata de N células entre portas de impedância escolhida. O botão **Exportar S2P** grava esses "
        "parâmetros S em formato Touchstone.")
    m.h("4.4  O que a célula modifica e quais os impactos", 2)
    m.tabela(["Ao aumentar…", "Efeito na linha periódica", "Impacto prático"],
             [["L série (Ls)", "Diminui a velocidade de fase, aumenta Z_B, baixa o corte",
               "Linha mais lenta e de maior impedância"],
              ["C em derivação (Cp)", "Diminui a velocidade de fase, reduz Z_B, baixa o corte",
               "Linha mais lenta e de menor impedância"],
              ["Cs série / Lp em derivação", "Introduz banda *left-handed* em baixa frequência (fase "
               "avança com a frequência)", "Base das linhas CRLH: dupla banda, antenas *leaky-wave*"],
              ["Número de células N", "Não altera dispersão nem Z_B; aprofunda a rejeição na banda "
               "proibida e aumenta o número de ondulações na banda passante", "Compromisso entre "
               "seletividade e tamanho"],
              ["Assimetria da célula", "Mesma dispersão; Z_B complexo", "Exige compensação reativa no "
               "casamento com a fonte e a carga"]],
             [3.4, 6.6, 6.0], "4.1", "Efeito dos elementos da célula concentrada.")
    cap4_banda_larga(m, F)


def cap4_banda_larga(m, F):
    m.h("4.5  Quando o objetivo é banda larga: redes passa-tudo", 2)
    m.p("Até aqui, as linhas periódicas tinham banda limitada por construção. Há três motivos: a "
        "periodicidade impõe um corte; a impedância de Bloch varia com a frequência; e o truncamento da "
        "estrutura gera ondulação na banda passante. Na escada LC da Eq. (4.4), Z_B cai de √(L/C) até zero no "
        "corte. Se o objetivo é que |S₂₁| e |S₁₁| se mantenham bons numa faixa larga, é preciso atacar "
        "esses três pontos.")
    m.p("**Regra prática para a escada.** O pior caso de reflexão de um trecho de linha com impedância Z_B "
        "entre portas Z₀ é |Γ| = |r² − 1|/(r² + 1), com r = Z_B/Z₀. Exigir |S₁₁| ≤ −15 dB (|Γ| ≤ 0,178) "
        "para qualquer comprimento leva a r ≥ 0,835, isto é, √(1 − (f/f_c)²) ≥ 0,835:")
    m.eq(r"f\leq0,55\,f_{c}\quad\Rightarrow\quad f_{c}\gtrsim1,8\,f_{\max}", "4.5")
    m.p("O LTA confirma a regra: com L = 2,5 nH e C = 1 pF (f_c = 6,37 GHz), a banda útil da escada vai até "
        "3,51 GHz = 0,55·f_c, para 5, 20 ou 50 células.")
    m.p("**Redes passa-tudo.** A saída mais radical está na seção 6.4 do livro (" + ref("6.4.1", "429–439") +
        "). Na **rede em treliça** (X), com impedância Z_s nos braços série e Z_p nas diagonais, a "
        "dispersão e a impedância são (Martín 6.13–6.14):")
    m.eq(r"\cos(\beta l)=\frac{Z_{p}+Z_{s}}{Z_{p}-Z_{s}},\qquad Z_{B}=\sqrt{Z_{s}Z_{p}}", "4.6")
    m.p("Com uma indutância nos braços e uma capacitância nas diagonais, Z_sZ_p = L/C é **real e "
        "independente da frequência**: a rede é passa-tudo, de resistência constante. A transmissão é "
        "total e só a fase varia, do DC até frequências arbitrariamente altas, e a mesma ideia vale para "
        "células CRLH de ordem 2, desde que polos e zeros dos dois ramos coincidam. O inconveniente é que a "
        "treliça é balanceada (diferencial). O livro mostra implementações com linhas acopladas, duas "
        "camadas metálicas ou vias (pp. 434–437). Uma alternativa de terminação simples é o **T em ponte**: "
        "braços série Z₁, derivação Z₂ e uma ponte Z₃ entre as portas. Com Z₃ = 4Z₂ (Martín 6.27–6.29):")
    m.eq(r"Z_{B}=\sqrt{2Z_{1}Z_{2}}", "4.7")
    m.p("e a rede é passa-tudo quando polos e zeros de Z₁ e Y₂ coincidem. O livro cita um protótipo em "
        "microstrip com banda de 1 a 8 GHz (p. 438). Na aba Célula ABCD, as topologias **Treliça (X)** e "
        "**T em ponte** usam os mesmos campos: ramo série = braços (Z₁), ramo paralelo = diagonais (Z₂). A "
        "ponte Z₃ = 4Z₂ é montada automaticamente. Os botões *Ex. passa-tudo (treliça)* e *Ex. T em ponte* "
        "carregam os casos das Figs. 6.25a e 6.31 do livro.")
    linhas = []
    for r in F["passa_tudo"]:
        fim = " (fim da faixa)" if r["ate_fim"] else ""
        linhas.append([r["nome"], f"{n(r['zb_min'], 1)} a {n(r['zb_max'], 1)} Ω", f"0 a {n(r['f_util'], 2)} GHz{fim}"])
    m.tabela(["Rede (20 células, portas de 50 Ω)", "Re Z_B na banda passante", "Banda útil (|S₁₁| ≤ −15 dB, |S₂₁| ≥ −1 dB)"],
             linhas, [5.4, 4.6, 6.0], "4.2", "Escada LC comparada às redes passa-tudo (cálculo até 20 GHz).")
    m.figura(IMG / "tela_13_trelica_resultados.png", "4.3",
             "Treliça (azul) contra escada T (laranja tracejado) com os mesmos L e C: Z_B constante e |S₂₁| = 0 dB "
             "até 20 GHz na treliça; a escada corta em 6,4 GHz. O |S₁₁| da treliça fica abaixo de −200 dB e não "
             "aparece no gráfico.")
    m.figura(IMG / "tela_15_ponte_estrutura.png", "4.4",
             "Esquemático do T em ponte gerado pelo LTA (C = 1 pF, L = 1,25 nH, ponte 4L = 5 nH): Z_B = 50 Ω "
             "constante.", 13)
    m.nota("as redes passa-tudo resolvem o casamento em banda larga, mas a fase continua variando com "
           "a frequência (são linhas dispersivas). Para atraso de grupo constante, a referência continua sendo "
           "a linha comum. Os elementos concentrados também têm parasitas que limitam a banda real, e a "
           "geometria dessas redes ainda não é gerada pelo LTA. Uma quarta via, os elementos *non-Foster* "
           "(capacitância negativa com circuitos ativos), é citada no livro com ressalvas de estabilidade "
           "(pp. 439–441).", "Limites")


# =============================================================================
def cap5(m, F):
    c, lc = F["cls"], F["livro_carga"]
    m.h("5  Linhas com carga capacitiva e o efeito slow-wave")
    m.p("A primeira aplicação concreta da teoria é carregar uma linha de transmissão com capacitâncias "
        "em derivação, espaçadas de um período l (" + ref("2.5.2", "102–105") + "). Essas capacitâncias "
        "podem ser componentes discretos ou, em tecnologia planar, pequenos trechos largos de trilha "
        "(*patches*), que têm muita capacitância para o terra. Aplicando a Eq. (4.2) à célula formada por "
        "um trecho de linha de impedância Z₀ e fase kl mais a capacitância C_ls, obtém-se a dispersão "
        "(Martín 2.116):")
    m.eq(r"\cos(\beta l)=\cos(kl)-\frac{\omega C_{ls}Z_{0}}{2}\sin(kl)", "5.1")
    m.p("em que k é a constante de fase da linha descarregada e β a da linha carregada. Quando C_ls é "
        "bem maior que a capacitância própria do trecho de linha, a célula pode ser modelada por elementos "
        "concentrados: L e C da linha (por célula) e C_ls em paralelo com C. Desse modelo saem fórmulas "
        "fechadas para as três grandezas de projeto:")
    m.eq(r"f_{c}=\frac{1}{\pi\sqrt{L(C+C_{ls})}},\qquad Z_{B}=\sqrt{\frac{L}{C+C_{ls}}},\qquad f_{s}=\frac{1}{2\sqrt{LC}}", "5.2")
    m.p("A frequência de corte f_c delimita a primeira banda passante; Z_B é a impedância de Bloch em "
        "baixa frequência; e f_s é o início da primeira banda espúria, que ocorre quando o trecho de linha "
        "descarregada tem meio comprimento de onda (Martín 2.117–2.120). A velocidade de fase em baixa "
        "frequência é reduzida pelo **fator slow-wave**:")
    m.eq(r"\frac{v_{p,\mathrm{sem\ carga}}}{v_{p,\mathrm{carregada}}}=\sqrt{\frac{C+C_{ls}}{C}}", "5.3")
    m.h("5.1  Projeto: das especificações aos elementos", 2)
    m.p("O livro mostra como, a partir das Eqs. (5.2), \"os valores dos elementos podem ser inferidos\". "
        "O LTA faz essa inversão de duas maneiras. Na primeira, o projetista especifica f_c, Z_B e f_s:")
    m.eq(r"L=\frac{Z_{B}}{\pi f_{c}},\qquad C+C_{ls}=\frac{1}{\pi f_{c}Z_{B}},\qquad C=\frac{1}{4f_{s}^{2}L}", "5.4")
    m.p("Na segunda, que é o procedimento usado no exemplo do livro (" + ref("2.5.2.2", "108–113") + "), "
        "especifica-se a fase φ que cada célula deve ter numa frequência f₀ (Martín 2.126), além de Z_B e f_s:")
    m.eq(r"\sqrt{L(C+C_{ls})}=\frac{\varphi}{2\pi f_{0}},\qquad L=Z_{B}\sqrt{L(C+C_{ls})},\qquad C+C_{ls}=\frac{\sqrt{L(C+C_{ls})}}{Z_{B}}", "5.5")
    m.p("e C novamente pela última das Eqs. (5.4). A especificação é inviável se resultar C_ls ≤ 0, o que "
        "acontece quando se pede f_s baixa demais para o corte desejado; o software avisa nesse caso.")
    m.h("5.2  Da célula elétrica à geometria microstrip", 2)
    m.p("Com L e C, a linha descarregada fica definida: Z₀ = √(L/C) e o atraso por célula √(LC). Na "
        "microstrip, a largura sai da síntese da Eq. (2.7) e o período vem do atraso:")
    m.eq(r"l=\frac{c\sqrt{LC}}{\sqrt{\varepsilon_{\mathrm{eff}}}}", "5.6")
    m.p("A capacitância de carga é implementada por um *patch*, um trecho largo de impedância Z_p "
        "(escolhida pelo projetista, tipicamente 10 a 20 Ω), centrado em cada célula. Como o patch "
        "substitui um pedaço da linha fina, o que conta é a capacitância **excedente**. Para um patch de "
        "comprimento l_p, a estimativa usada pelo LTA é:")
    m.eq(r"C_{ls}\approx\frac{l_{p}}{c}\left(\frac{\sqrt{\varepsilon_{\mathrm{eff},p}}}{Z_{p}}-\frac{\sqrt{\varepsilon_{\mathrm{eff}}}}{Z_{0}}\right)", "5.7")
    m.nota("a Eq. (5.7) é uma aproximação quase-estática que ignora os efeitos de borda e de "
           "degrau entre a linha fina e o patch. Ela fornece um **ponto de partida** para a otimização em "
           "simulador eletromagnético, exatamente como foi feito no livro com o Agilent Momentum. Se o "
           "patch calculado não couber no período, reduza Z_p (patch mais largo e mais curto).", "Atenção")
    m.h("5.3  Exemplo do livro: filtro UWB com supressão de espúrios", 2)
    m.p("No exemplo do livro, os trechos de linha de 70° e 47,5 Ω entre os *stubs* de um filtro passa-faixa "
        "de banda larga são substituídos por duas células carregadas com patches. A especificação é de 35° "
        "por célula a 1,4 GHz, Z_B = 47,5 Ω e f_s = 22 GHz, para empurrar a primeira espúria para bem longe. "
        "Substituindo nas Eqs. (5.5) e (5.4):")
    m.tabela(["Grandeza", "LTA (Eqs. 5.4–5.5)", "Livro (Fig. 2.34)"],
             [["L por célula", f"{n(lc['L'], 2)} nH", "3,1 nH"],
              ["C por célula", f"{n(lc['C'], 3)} pF", "0,17 pF"],
              ["C_ls", f"{n(lc['Cls'], 2)} pF", "1,2 pF"],
              ["Z₀ da linha descarregada", f"{n(lc['zo'], 1)} Ω", "123,5 Ω"],
              ["Frequência de corte f_c", f"{n(lc['fc'], 2)} GHz", "5,6 GHz"],
              ["Fator slow-wave", f"{n(lc['sw'], 2)}", "–"]],
             [5.4, 5.3, 5.3], "5.1", "Elementos do exemplo do livro (substrato Arlon, εr = 2,4, h = 0,675 mm).")
    m.p("A diferença de cerca de 10% nos elementos se explica porque os valores do livro já passaram por "
        "otimização e arredondamento, e porque o livro fixou f_c em 5,6 GHz em vez de calculá-la a partir "
        "da fase. A ordem de grandeza e o comportamento são os mesmos. No layout, o LTA propõe um período de "
        f"{n(lc['periodo'], 2)} mm, linha de {n(lc['w_linha'], 3)} mm e patch de {n(lc['w_patch'], 2)} × "
        f"{n(lc['l_patch'], 2)} mm (com Z_p = 10 Ω).")
    m.figura(IMG / "tela_04_carga_estrutura.png", "5.1",
             "Aba Carga capacitiva, sub-aba Estrutura, logo após mudar Z_p de 10 para 15 Ω. A versão "
             "anterior aparece tracejada; as cotas (a–e) estão no quadro; o corte A mostra o patch sobre o "
             "substrato; e a caixa inferior lista as dimensões que mudaram.")
    m.figura(IMG / "tela_05_carga_resultados.png", "5.2",
             "Resultados: diagrama de Brillouin (modelo distribuído da Eq. 5.1 e concentrado), impedância de "
             "Bloch, parâmetros S de duas células e defasagem por célula na primeira banda.")
    m.p("Repare, no diagrama de Brillouin, que as curvas distribuída e concentrada praticamente coincidem "
        "até o corte, e que a curva pontilhada da linha sem carga é muito mais íngreme: essa diferença de "
        "inclinação é o efeito *slow-wave*. A primeira banda espúria só aparece perto de 22 GHz, como "
        "especificado.")
    m.h("5.4  O que a carga capacitiva modifica e quais os impactos", 2)
    m.p("A Figura 5.3 isola o efeito da capacitância de carga. A linha descarregada é mantida fixa "
        f"(L = {n(c['L'] * 1e9, 1)} nH e C = {n(c['C'] * 1e12, 1)} pF por célula, Z₀ = {n(c['zo'], 0)} Ω) "
        "e C_ls é variada.")
    m.figura(IMG / "fig_impacto_cls.png", "5.3",
             "Efeito da razão C_ls/C (calculado com o LTA). (a) A curva de dispersão fica mais deitada e o "
             "corte desce. (b) A impedância de Bloch cai e passa a variar com a frequência perto do corte.")
    linhas = []
    for r, zb, fc, sw in c["tabela"]:
        linhas.append([n(r, 0), n(zb, 1), "–" if r == 0 else n(fc, 2), n(sw, 2)])
    m.tabela(["C_ls / C", "Z_B [Ω]", "f_c [GHz]", "Fator slow-wave"], linhas, [4, 4, 4, 4], "5.2",
             "Valores correspondentes à Figura 5.3.")
    m.p("Em resumo, carregar a linha com capacitância produz, ao mesmo tempo, três efeitos que o projetista "
        "precisa equilibrar:")
    m.lista(["**Miniaturização:** a onda fica √((C + C_ls)/C) vezes mais lenta, e todo trecho de fase "
             "especificada fica proporcionalmente mais curto. É uma alternativa a substratos de εr alto.",
             "**Queda de impedância:** Z_B = √(L/(C + C_ls)) cai. Para manter 50 Ω, a linha descarregada "
             "precisa ter Z₀ bem acima de 50 Ω (trilha fina). Na prática, a largura mínima fabricável "
             "limita o quanto se pode carregar.",
             "**Filtragem intrínseca:** surge um corte em f_c e a primeira espúria é empurrada para f_s. "
             "Escolhendo f_c entre a banda útil e a primeira espúria de um filtro distribuído, elimina-se "
             "a espúria sem componentes adicionais."])
    m.h("5.5  Aplicações descritas na literatura", 2)
    m.lista(["**Filtros CPW compactos com supressão de espúrios:** aproximando periodicamente a fita "
             "central dos planos de terra, um filtro passa-faixa de 4ª ordem a 6 GHz ficou cerca de 30% "
             "menor e livre de espúrias numa banda larga (" + ref("2.5.2.1", "105–108") + ").",
             "**Filtro UWB em microstrip:** a substituição dos trechos entre *stubs* por linhas carregadas "
             "com patches deu cerca de 50% de redução de comprimento e rejeição melhor que −30 dB até "
             "20 GHz, com perdas na banda abaixo de 0,9 dB (" + ref("2.5.2.2", "108–113") + ").",
             "**Híbridas e acopladores miniaturizados**, e linhas *slow-wave* em CMOS para ondas "
             "milimétricas, citados no livro como extensões do mesmo conceito (p. 108)."])


# =============================================================================
def cap6(m, F):
    le, e = F["livro_ebg"], F["ebg"]
    m.h("6  Linhas EBG e a teoria de modos acoplados")
    m.p("Uma segunda família de linhas periódicas não usa elementos concentrados: a própria geometria "
        "varia ao longo do comprimento. Pode-se modular a largura da trilha (o que desloca a impedância "
        "local para cima e para baixo) ou gravar furos no plano de terra sob a trilha. Essas estruturas são "
        "chamadas de **EBG** (*electromagnetic bandgap*), em analogia com os cristais fotônicos da óptica, "
        "e servem sobretudo para **rejeitar** uma faixa de frequências (" + ref("2.1", "47–48") + ").")
    m.p("Para analisá-las, o livro usa a **teoria de modos acoplados** (" + ref("2.4.1", "65–69") + "). A "
        "ideia é olhar cada seção transversal da linha não uniforme como se fosse uma linha uniforme "
        "\"auxiliar\" com aquela mesma seção. O campo total é escrito como a soma de uma onda direta a⁺(z) e "
        "uma onda refletida a⁻(z) dessa linha auxiliar, cujas amplitudes variam com z. Onde a geometria muda, "
        "parte da onda direta é transferida para a refletida. Com as aproximações de modo único e sem "
        "radiação, as equações ficam (Martín 2.48):")
    m.eq(r"\frac{d}{dz}\begin{bmatrix}a^{+}\\a^{-}\end{bmatrix}=\begin{bmatrix}-j\beta&K\\K&j\beta\end{bmatrix}\begin{bmatrix}a^{+}\\a^{-}\end{bmatrix}", "6.1")
    m.p("O acoplamento entre as duas ondas é medido pelo **coeficiente de acoplamento** K(z), que "
        "depende apenas de quão rápido a impedância característica local varia (Martín 2.50, deduzida no "
        "Apêndice E do livro, pp. 482–483):")
    m.eq(r"K(z)=-\frac{1}{2Z_{0}(z)}\frac{dZ_{0}(z)}{dz}", "6.2")
    m.p("Essa expressão é a chave do projeto: **onde a impedância não varia, não há acoplamento**; onde ela "
        "sobe ou desce rapidamente, há transferência de energia para a onda refletida. Resolvidas as "
        "equações com uma onda incidente unitária na porta 1 e a porta 2 casada, os parâmetros S saem "
        "diretamente: S₁₁ = a⁻(0)/a⁺(0) e S₂₁ = a⁺(L)/a⁺(0) (" + ref("2.4.2", "69–70") + ").")
    m.h("6.1  Casamento de fase e solução analítica", 2)
    m.p("Se a perturbação é periódica, K(z) também é, e pode ser expandido em série de Fourier com "
        "coeficientes K_n. Perto da frequência em que β = nπ/l, apenas o termo K_n contribui de forma "
        "acumulada ao longo de muitos períodos (os demais oscilam e se cancelam). Essa é a **condição de "
        "casamento de fase**, a mesma condição de Bragg da Eq. (3.3), agora deduzida matematicamente "
        "(" + ref("2.4.3", "71–77") + "). Definindo o desvio de fase e a constante de decaimento")
    m.eq(r"\Delta\beta=\beta-\frac{n\pi}{l},\qquad\gamma=\sqrt{|K_{n}|^{2}-\Delta\beta^{2}}", "6.3")
    m.p("as equações têm solução fechada para uma estrutura de comprimento L (Martín 2.76):")
    m.eq(r"S_{11}=\frac{|K_{n}|\sinh(\gamma L)}{-\Delta\beta\sinh(\gamma L)+j\gamma\cosh(\gamma L)},\qquad S_{21}=\frac{j\gamma\,e^{-jn\pi L/l}}{-\Delta\beta\sinh(\gamma L)+j\gamma\cosh(\gamma L)}", "6.4")
    m.p("Quando |Δβ| < |K_n|, γ é real e a onda decai exponencialmente: é a banda rejeitada. Daí saem as "
        "expressões de projeto mais úteis do capítulo (" + ref("2.4.4", "77–79") + "):")
    m.eq(r"f_{\max}=\frac{nc}{2l\sqrt{\bar{\varepsilon}_{re}}},\qquad|S_{11}|_{\max}=\tanh(|K_{n}|L),\qquad|S_{21}|_{\min}=\mathrm{sech}(|K_{n}|L)", "6.5")
    m.eq(r"BW=\frac{c|K_{n}|}{\pi\sqrt{\bar{\varepsilon}_{re}}}\sqrt{1+\left(\frac{\pi}{|K_{n}|L}\right)^{2}}", "6.6")
    m.p("em que ε̄_re é a permissividade efetiva média ao longo do período, calculada como "
        "√ε̄_re = (1/l)∫√ε_re(z)dz (Martín 2.57, justificada no Apêndice F, pp. 484–485). A Eq. (6.5) "
        "mostra que **a rejeição depende só do produto |K_n|·L**: perturbação mais forte ou estrutura mais "
        "longa produzem o mesmo efeito.")
    m.h("6.2  O conteúdo harmônico de K(z) decide quantas bandas aparecem", 2)
    m.p("Em condições de baixa refletividade, o livro mostra que S₁₁ é aproximadamente a transformada de "
        "Fourier do coeficiente de acoplamento (" + ref("2.4.5", "79–80") + "):")
    m.eq(r"S_{11}\approx-\int_{0}^{L}K(z)\,e^{-j2\beta z}\,dz", "6.7")
    m.p("A consequência é uma regra de projeto muito clara. Se K(z) for **senoidal**, há um único "
        "coeficiente de Fourier e, portanto, uma única banda rejeitada. Se K(z) for uma **onda quadrada**, "
        "aparecem bandas nas harmônicas ímpares, com |K_n| = 2K/(nπ) e rejeição cada vez menor. Os perfis "
        "de impedância correspondentes são obtidos integrando a Eq. (6.2) (Martín 2.102 e 2.104):")
    m.eq(r"Z_{0}(z)=Z_{\mathrm{ref}}\exp\left[\frac{2|K_{1}|l}{\pi}\sin\left(\frac{2\pi z}{l}\right)\right]\quad\text{(senoidal)}", "6.8")
    m.p("Uma consequência útil da Eq. (6.8), que usaremos no capítulo 8: a razão máxima de impedância é "
        "Z_max/Z_ref = exp(2|K₁|l/π). Fixada essa razão (limitada pela largura de trilha fabricável), o "
        "produto |K₁|·l fica fixo, e o número de períodos necessário para uma rejeição desejada passa a "
        "não depender da frequência:")
    m.eq(r"N\geq\frac{\mathrm{arccosh}\left(10^{-S_{21,\mathrm{dB}}/20}\right)}{\frac{\pi}{2}\ln\left(Z_{\max}/Z_{\mathrm{ref}}\right)}", "6.9")
    m.h("6.3  Solução numérica e verificação no LTA", 2)
    m.p("O LTA resolve a Eq. (6.1) numericamente, sem as aproximações da Eq. (6.4). A linha é dividida em "
        "segmentos curtos, em cada um dos quais β e K são constantes; a matriz de transição de cada "
        "segmento é a exponencial exata exp(MΔz) = cosh(gΔz)·I + sinh(gΔz)/g·M, com g² = K² − β². O "
        "produto dessas matrizes dá a relação entre as ondas nas duas pontas, e as condições de contorno "
        "fornecem S₁₁, S₂₁ e S₂₂. Como verificação independente, o software calcula também a cascata de "
        "matrizes ABCD de trechos uniformes curtos. As duas curvas devem coincidir, e a solução analítica "
        "da Eq. (6.4) é sobreposta perto de cada banda.")
    m.p(f"Para o exemplo da Fig. 2.11 do livro (RO3010, l = 10 mm, 10 períodos, |K₁| = 0,64 cm⁻¹, perfil "
        f"senoidal), o LTA obtém |S₂₁|_min = {n(le['s21'], 1)} dB (livro: −49,6 dB), banda de "
        f"{n(le['bw'], 2)} GHz (livro: 2,55 GHz) e f_max = {n(le['fmax'], 2)} GHz (livro: 5,66 GHz). A "
        f"pequena diferença em f_max vem de ε̄_re = {n(le['eem'], 2)}, calculada por fórmula fechada, enquanto "
        "o livro usou solver eletromagnético.")
    m.figura(IMG / "tela_06_ebg_sen_resultados.png", "6.1",
             "Aba EBG / Modos acoplados com o exemplo senoidal do livro: perfil Z₀(z) e W(z), |S₂₁| e |S₁₁| "
             "pelos três métodos (CMT numérico, cascata ABCD e analítico) e coeficiente de acoplamento K(z).")
    m.figura(IMG / "tela_08_ebg_quad_resultados.png", "6.2",
             "Perfil quadrado (Fig. 2.14 do livro, K = 1 cm⁻¹): bandas em f_max e 3f_max, como previsto "
             "pelo conteúdo harmônico de K(z). A tabela à esquerda dá −49,3, −12,5 e −5,7 dB para "
             "n = 1, 3 e 5, idênticos aos valores do livro.")
    m.figura(IMG / "tela_07_ebg_quad_estrutura.png", "6.3",
             "Pré-visualização da estrutura quadrada logo após reduzir |K| de 1,0 para 0,6 cm⁻¹. A trilha "
             "anterior aparece tracejada, e a caixa inferior mostra que a faixa de impedância caiu de "
             "30–82 Ω para 37–68 Ω.")
    m.h("6.4  Furos no plano de terra", 2)
    m.p("A outra forma clássica de perturbação é gravar furos circulares no plano de terra sob a trilha "
        "(Martín, Fig. 2.15). O livro mostra que uma única fileira de furos, alinhada com a trilha, tem "
        "praticamente a mesma resposta que a matriz bidimensional original, com grande economia de área "
        "(pp. 86–88). A frequência da banda rejeitada continua dada pela condição de Bragg, mas a impedância "
        "local Z₀(z) sobre um furo não tem fórmula fechada: exige solução eletromagnética (Martín 2.51). Por "
        "isso o LTA gera a geometria e estima f_Bragg, deixando a resposta em frequência para o simulador.")
    m.figura(IMG / "tela_09_ebg_furos_estrutura.png", "6.4",
             "EBG com furos no terra (exemplo da Fig. 2.16 do livro: l = 18,9 mm, r/l = 0,25, RO3010). O "
             "corte A passa pelo centro do primeiro furo e mostra a abertura no plano de terra.")
    m.h("6.5  O que a perturbação modifica e quais os impactos", 2)
    m.figura(IMG / "fig_impacto_ebg.png", "6.5",
             "Efeito dos parâmetros de uma EBG senoidal em RO3010 com l = 10 mm (CMT numérico, LTA). "
             "(a) Aumentar |K| aprofunda e alarga a banda. (b) Aumentar o número de períodos aprofunda a "
             "banda e a deixa mais estreita, até um limite.")
    linhas = [[n(k, 2), f"{n(zmin, 0)}–{n(zmax, 0)}", n(s, 1), n(bw, 2)] for k, s, bw, zmin, zmax in e["k"]]
    m.tabela(["|K| [cm⁻¹]", "Faixa de Z₀ [Ω]", "|S₂₁|min [dB]", "Banda [GHz]"], linhas, [4, 4, 4, 4], "6.1",
             "Efeito da amplitude da perturbação (N = 10, Figura 6.5a).")
    linhas = [[str(nn), n(s, 1), n(bw, 2)] for nn, s, bw in e["n"]]
    m.tabela(["Períodos N", "|S₂₁|min [dB]", "Banda [GHz]"], linhas, [5.3, 5.3, 5.4], "6.2",
             "Efeito do número de períodos (|K| = 0,4 cm⁻¹, Figura 6.5b).")
    m.tabela(["Parâmetro", "O que muda", "Compromisso"],
             [["Período l", "Posição da banda: f_max ∝ 1/l (Eq. 6.5)",
               "Frequências baixas exigem períodos longos e estruturas grandes"],
              ["Amplitude |K| (variação de W ou tamanho do furo)", "Profundidade e largura da banda",
               "Trilhas muito finas ou largas; mais ondulação na banda passante; mais radiação"],
              ["Número de períodos N", "Profundidade (tanh/sech de |K|L)", "Comprimento total"],
              ["Forma de K(z)", "Quais harmônicas são rejeitadas (Eq. 6.7)",
               "Senoidal: banda única; quadrada: bandas ímpares; soma de senoides: bandas escolhidas"],
              ["Janelamento / *tapering* nas pontas", "Reduz a ondulação na banda passante",
               "Aumenta o comprimento para a mesma rejeição"],
              ["*Chirp* (período variando ao longo da linha)", "Alarga a banda rejeitada; atraso de grupo "
               "linear com a frequência", "Base das linhas de atraso chirpadas"]],
             [4.4, 5.6, 6.0], "6.3", "Parâmetros de uma linha EBG e seus efeitos.")
    m.p("**Banda passante larga numa EBG.** A ondulação na banda passante limita a faixa útil abaixo da "
        "banda rejeitada. No perfil senoidal do livro, |S₁₁| ultrapassa −15 dB bem antes do início da "
        "rejeição. Multiplicar K(z) por uma janela, o que o perfil personalizado da seção 6.7 permite com "
        "`janela(z)` ou `hann(z)`, suaviza as pontas e amplia essa faixa:")
    linhas = [[r["nome"], f"`{r['expr']}`", f"{n(r['zmin'], 0)}–{n(r['zmax'], 0)}", n(r["s21"], 1),
               f"até {n(r['f_util'], 2)} GHz"] for r in F["janela"]]
    m.tabela(["Perfil", "K(z) [1/cm]", "Z₀ [Ω]", "|S₂₁|min [dB]", "Banda útil inferior"], linhas,
             [3.4, 5.0, 2.0, 2.4, 3.2], "6.3a",
             "Efeito do janelamento (RO3010, l = 10 mm, critério |S₁₁| ≤ −15 dB e |S₂₁| ≥ −1 dB).", fonte=8.5)
    m.h("6.6  Aplicações descritas na literatura", 2)
    m.lista(["**Refletores e filtros rejeita-faixa:** furos no terra ou modulação da trilha; *tapering* "
             "para reduzir ondulação, *chirping* para alargar a banda, e linhas dobradas para reduzir o "
             "comprimento (" + ref("2.5.1.1", "86–92") + ").",
             "**Ressonadores de Bragg de alto Q:** dois refletores EBG separados por um trecho de linha, com "
             "Q medido de 129 a 4,2 GHz, valor alto para estrutura planar (" + ref("2.5.1.2", "92–93") + ").",
             "**Supressão de espúrios em filtros de linhas acopladas** (filtro *wiggly-line*): modulação "
             "senoidal da borda das linhas, com período λg/4, rejeitando a primeira banda espúria em mais de "
             "40 dB sem alterar a banda útil (" + ref("2.5.1.3", "93–95") + ").",
             "**Supressão de harmônicos em circuitos ativos:** numa antena ativa (oscilador a 4,5 GHz "
             "alimentando um *patch*), uma EBG dupla sintonizada em 2f₀ e 3f₀ reduziu esses harmônicos em "
             "cerca de 20 e 10 dB e aumentou a potência na fundamental em cerca de 2,6 dB, sem aumentar a "
             "área (" + ref("2.5.1.4", "95–99") + ").",
             "**Linhas de atraso chirpadas** para processamento analógico de sinais: atraso de grupo linear "
             "com a frequência numa banda de 12 GHz (" + ref("2.5.1.5", "99–102") + ")."])
    cap6_personalizado(m, F)


def cap6_personalizado(m, F):
    m.h("6.7  Projetando seus próprios perfis: bandas em várias frequências", 2)
    m.p("Os perfis senoidal e quadrado são casos particulares. O livro mostra EBGs sintonizadas em duas e "
        "três frequências, obtidas somando senoides na geometria do plano de terra, com janela de Hamming "
        "para reduzir a ondulação (" + ref("2.5.1.1", "88–91") + "), e a antena ativa com EBG dupla em 2f₀ e "
        "3f₀ (pp. 95–99). A Eq. (6.7) explica por que isso funciona: como S₁₁ é aproximadamente a "
        "transformada de Fourier de K(z), **cada componente senoidal do coeficiente de acoplamento cria a "
        "sua própria banda rejeitada**. Para posicionar uma banda em f_i, basta que o período do termo "
        "correspondente satisfaça a condição de Bragg:")
    m.eq(r"K(z)=\sum_{i}A_{i}\cos\left(\frac{2\pi z}{l_{i}}\right),\qquad l_{i}=\frac{c}{2f_{i}\sqrt{\varepsilon_{\mathrm{eff}}}}", "6.10")
    m.p("Cada termo tem coeficiente de Fourier |K_n| = A_i/2. Enquanto as bandas estiverem bem separadas, "
        "a profundidade de cada uma segue a Eq. (6.5):")
    m.eq(r"|S_{21}|_{\min,i}\approx\mathrm{sech}\left(\frac{A_{i}L}{2}\right)", "6.11")
    m.p("Com A = 0,4 cm⁻¹ e L = 15 cm, por exemplo, A·L/2 = 3 e |S₂₁|_min ≈ sech(3), cerca de −20 dB, o que "
        "a solução numérica confirma. A impedância local é obtida integrando a Eq. (6.2), "
        "Z₀(z) = Z_ref·exp(−2∫K dz) (Martín 2.100). O LTA remove o valor médio de K antes da integração, "
        "para que a linha termine na mesma impedância em que começa e não haja degrau na porta 2.")
    m.h("Como usar no LTA", 3)
    m.p("Na aba EBG / Modos acoplados, escolha a perturbação **Personalizado (função)**. O painel *Perfil "
        "personalizado* aceita uma expressão em z (em mm, de 0 a L) que pode descrever K(z) em cm⁻¹, Z₀(z) "
        "em Ω ou W(z) em mm. A pré-visualização é atualizada enquanto se digita, e o botão *Carregar* "
        "preenche um dos modelos prontos. A expressão é verificada antes de ser avaliada: só números, "
        "operadores, comparações e as funções da Tabela 6.4 são aceitos.")
    m.tabela(["Nome", "Significado"],
             [["`z`, `L`, `Zref`", "Posição (mm), comprimento total (mm) e impedância de referência (Ω)"],
              ["`bragg(f)`", "Período em mm cuja primeira banda cai em f GHz (Eq. 3.3), usando o εeff da linha "
               "de referência"],
              ["`janela(z)`, `hann(z)`", "Janelas de Hamming e de Hann em [0, L], para reduzir a ondulação"],
              ["`ret(z, a, b)`, `degrau(z, a)`", "Pulso retangular entre a e b; degrau a partir de a"],
              ["`sin`, `cos`, `exp`, `log`, `sqrt`, `tanh`, `abs`, `sign`, `where`, `clip`, `pi`",
               "Funções matemáticas usuais (NumPy)"]],
             [5.0, 11.0], "6.4", "Variáveis e funções disponíveis na expressão do perfil personalizado.")
    m.figura(IMG / "tela_11_ebg_pers_estrutura.png", "6.6",
             "Perfil personalizado com duas bandas, logo após mudar a segunda de 5 para 5,5 GHz na expressão. "
             "A versão anterior aparece tracejada; o detalhe mostra o trecho de maior largura.")
    m.figura(IMG / "tela_12_ebg_pers_resultados.png", "6.7",
             "Resultados do perfil de duas bandas: Z₀(z) e W(z), bandas em cerca de 3 e 5,5 GHz pelos "
             "métodos CMT e ABCD, estimativa de Fourier de K(z) sobre o |S₁₁| (pontilhado) e o próprio K(z).")
    m.p("A curva pontilhada de |S₁₁| é a estimativa da Eq. (6.7). Ela é imediata de calcular e mostra, antes "
        "mesmo da solução completa, onde o conteúdo espectral de K(z) vai produzir reflexão. Perto das bandas "
        "a estimativa satura em 0 dB, porque só vale para baixa refletividade; ali vale a solução numérica. "
        "As bandas abaixo de −10 dB são detectadas automaticamente e listadas no painel de resultados.")
    linhas = []
    for r in F["pers"]:
        bandas = "; ".join(f"{n(fc, 2)} GHz ({n(db, 0)} dB)" for fc, db in r["bandas"][:4]) or "–"
        linhas.append([r["nome"], f"`{r['expr']}`", f"{n(r['comp'], 0)} mm",
                       f"{n(r['wmin'], 2)}–{n(r['wmax'], 2)}", bandas])
    m.tabela(["Modelo", "Expressão", "L", "W [mm]", "Bandas detectadas"], linhas,
             [3.2, 6.0, 1.4, 1.8, 3.6], "6.5",
             "Modelos prontos do LTA em RO3010 (h = 1,27 mm), 50 Ω, calculados por CMT até 18 GHz.", fonte=8)
    m.nota("as bandas reais usam a permissividade efetiva **média** do perfil, ligeiramente diferente "
           "da usada em `bragg(f)`; desvios de 1 a 2% são normais e podem ser corrigidos ajustando f na "
           "expressão. Termos muito próximos interagem, e amplitudes grandes produzem trilhas muito finas ou "
           "largas: confira sempre W mínimo e máximo na pré-visualização. Como nos demais casos, a geometria "
           "deve ser confirmada em simulador eletromagnético.", "Cuidados")


# =============================================================================
def cap7(m):
    m.h("7  Fluxo de trabalho no LTA e exportação")
    m.h("7.0  Navegação", 2)
    m.p("A janela mostra uma ferramenta por vez, escolhida pela barra de menus. Em **Ferramentas**, as "
        "ferramentas estão agrupadas por capítulo do livro (por enquanto, *Cap. 2 – Linhas periódicas*: "
        "Célula ABCD, Carga capacitiva e EBG / Modos acoplados). **Arquivo** dá acesso ao substrato e à pasta "
        "de saídas; **Exportar** abre a ferramenta de exportação; **Ajuda** abre este material, o manual "
        "técnico e a lista de atalhos. Abaixo do menu ficam a barra do **substrato ativo**, sempre visível e "
        "com o botão *Alterar…*, e o título da ferramenta atual. Os atalhos **Ctrl+1**, **Ctrl+2** e "
        "**Ctrl+3** alternam entre as ferramentas do capítulo 2, e **Ctrl+E** abre a exportação. Cada "
        "ferramenta guarda seus valores ao trocar de tela.")
    m.p("O fluxo sugerido é:")
    m.lista(["**Substrato:** escolher ou digitar εr, h, espessura de cobre e tangente de perdas. Todas as "
             "ferramentas passam a usar esses valores.",
             "**Especificar** a estrutura na aba correspondente (Célula ABCD, Carga capacitiva ou EBG). Os "
             "botões *Exemplo* carregam os casos do livro e são um bom ponto de partida.",
             "**Observar a pré-visualização** enquanto edita: a sub-aba *Estrutura* é redesenhada cerca de "
             "0,4 s depois que você para de digitar. Use *Fixar referência* para comparar várias alterações "
             "contra uma versão de partida.",
             "**Calcular** para ver dispersão, impedância de Bloch e parâmetros S na sub-aba *Resultados*.",
             "**Enviar para exportação** e, na aba *Exportar*, gerar os arquivos para a ferramenta desejada.",
             "**Simular em EM** (HFSS, CST, ADS Momentum), otimizar as dimensões e só então fabricar e medir."],
            numerada=True)
    m.h("7.1  A pré-visualização da estrutura", 2)
    m.p("A pré-visualização tem três vistas: a placa inteira, com dimensões totais e a célula do zoom "
        "destacada; o zoom da célula unitária, com cotas identificadas por letras e o quadro de valores; e o "
        "corte transversal A, com substrato, plano de terra (e abertura do furo, se houver) e trilha. A "
        "espessura do cobre é exagerada nessa vista para ficar visível. A versão anterior aparece tracejada e "
        "a caixa inferior lista as alterações. Especificações inviáveis aparecem em vermelho sobre a última "
        "estrutura válida, e trilhas mais finas que cerca de 0,1 mm geram um alerta de fabricabilidade.")
    m.h("7.2  Exportação", 2)
    m.figura(IMG / "tela_10_exportar.png", "7.1",
             "Aba Exportar: seleção da estrutura, formatos de layout, faixa de simulação, pasta de saída "
             "(relativa ao projeto) e exportação dos parâmetros S.")
    m.tabela(["Arquivo", "Ferramenta", "Como abrir"],
             [["`.kicad_pcb`", "KiCad 6/7/8; Altium", "Abrir no editor de PCB e pressionar B para preencher "
               "o terra. No Altium: File › Import Wizard › KiCad"],
              ["`.dxf` (R12)", "Altium, ADS, HFSS, CST", "Importar mapeando as camadas `TOPO_COBRE`, `TERRA`, "
               "`TERRA_FUROS` e `PLACA`"],
              ["`.gtl` / `.gbl` / `.gko`", "Altium (CAMtastic), fabricação", "Gerber RS-274X em mm"],
              ["`.gds`", "ADS, HFSS, CST", "GDSII; camadas 1 (topo), 2 (terra), 3 (furos), 10 (placa)"],
              ["`_hfss.py`", "Ansys HFSS / AEDT", "Tools › Run Script; cria projeto, variáveis `h_sub` e "
               "`h_ar`, portas de onda, radiação, setup e varredura"],
              ["`_cst.bas`", "CST Studio Suite", "Macro VBA; parâmetros h, t, er e tand editáveis"],
              ["`.s2p` / `_S.csv`", "ADS, AWR, planilhas", "Parâmetros S do modelo de circuito (Touchstone "
               "v1, RI, GHz) e CSV com dB e fase"]],
             [3.2, 3.8, 9.0], "7.1", "Formatos exportados pelo LTA.")
    m.p("Os parâmetros S exportados vêm dos **modelos ideais** usados nos cálculos (sem perdas, sem "
        "descontinuidades e sem radiação). Eles são úteis para estudos de sistema, por exemplo carregando o "
        ".s2p num bloco S2P do ADS e cascateando com o restante do circuito, mas não substituem o resultado "
        "de simulação eletromagnética ou de medição.")
    m.nota("os scripts para HFSS e CST e o arquivo do KiCad seguem a sintaxe documentada dessas "
           "ferramentas, mas ainda não foram executados nas versões instaladas no Instituto. Recomenda-se "
           "testar cada formato com um dos exemplos antes do uso em projeto.", "Situação da exportação")


# =============================================================================
def cap8(m, F):
    r, h, mi = F["refletor"], F["harmonico"], F["mini"]
    m.h("8  Exemplos de aplicação resolvidos")
    m.p("Os três exemplos a seguir mostram como os conceitos da aula se traduzem em decisões de projeto. "
        "Todos os números foram calculados com o núcleo do LTA e podem ser reproduzidos nas abas "
        "correspondentes.")
    # ---- refletor
    m.h("8.1  Refletor (rejeita-faixa) a 2,45 GHz em FR-4", 2)
    m.p("**Especificação.** Rejeitar a faixa ISM de 2,45 GHz com |S₂₁| ≤ −30 dB no centro, numa linha de "
        "50 Ω em FR-4 (εr = 4,4, h = 1,6 mm). Por fabricação, a impedância local deve ficar entre cerca de "
        "33 e 75 Ω, ou seja, Z_max/Z_ref = 1,5.")
    m.p(f"**Passo 1: número de períodos.** Pela Eq. (6.5), |S₂₁|_min = −30 dB exige |K|·L ≥ "
        f"arccosh(31,6) = {n(r['kl_min'], 2)}. Com Z_max/Z_ref = 1,5, cada período contribui "
        f"|K|·l = (π/2)·ln 1,5 = {n(r['kl_periodo'], 3)}. Logo, pela Eq. (6.9), N ≥ "
        f"{n(r['kl_min'] / r['kl_periodo'], 2)}, isto é, **N = {r['n']} períodos**.")
    m.p(f"**Passo 2: período.** A linha de 50 Ω tem W = {n(r['w50'], 2)} mm e εeff = {n(r['ee50'], 2)}. Pela "
        f"condição de Bragg, l ≈ c/(2f√εeff) = {n(r['l0'], 1)} mm. Refinando com a permissividade média do "
        f"perfil modulado (ε̄_re = {n(r['eem'], 3)}), obtém-se **l = {n(r['l'], 2)} mm** e "
        f"|K| = {n(r['k'], 3)} cm⁻¹.")
    m.p(f"**Resultado.** Comprimento total de {n(r['comp'], 0)} mm, trilha variando entre "
        f"{n(r['wmin'], 2)} e {n(r['wmax'], 2)} mm, f_max = {n(r['fmax'], 2)} GHz, "
        f"|S₂₁|_min = {n(r['s21'], 1)} dB (analítico) e {n(r['s21_num'], 1)} dB (CMT numérico), "
        f"banda de {n(r['bw'], 2)} GHz.")
    m.figura(IMG / "fig_ex_refletor.png", "8.1", "Resposta do refletor projetado (CMT numérico).", 12)
    m.p("**Discussão.** O comprimento de quase um quarto de metro é a lição do exemplo: EBGs em frequências "
        "baixas são longas. As saídas são aceitar Z_max/Z_ref maior (menos períodos, trilhas mais "
        "extremas), dobrar a linha em meandro, como na Fig. 2.18 do livro, ou usar furos no terra, que "
        "permitem perturbação mais forte. Para as faixas de 5–6 GHz, o mesmo projeto fica com menos da "
        "metade do comprimento.")
    # ---- harmonico
    m.h("8.2  Supressão do segundo harmônico de um oscilador", 2)
    m.p("**Especificação.** Um oscilador de 2,45 GHz em RO4003C (εr = 3,55, h = 0,813 mm) tem segundo "
        "harmônico forte em 4,9 GHz. Deseja-se uma seção EBG na saída que atenue 2f₀ sem afetar f₀, nos "
        "moldes do exemplo da antena ativa do livro (pp. 95–99).")
    m.p(f"**Projeto.** Perfil senoidal (banda única, Eq. 6.7, para não atenuar a fundamental) com "
        f"Z_max/Z_ref = 1,5 e N = {h['n']} períodos, centrado em 2f₀. Resulta l = {n(h['l'], 2)} mm, "
        f"|K| = {n(h['k'], 3)} cm⁻¹ e comprimento de {n(h['comp'], 0)} mm.")
    m.p(f"**Resultado.** Em 2f₀: |S₂₁| = {n(h['s21_2f0'], 1)} dB. Em f₀: |S₂₁| = {n(h['s21_f0'], 2)} dB e "
        f"|S₁₁| = {n(h['s11_f0'], 1)} dB. A banda rejeitada tem {n(h['bw'], 2)} GHz e fica longe da "
        "fundamental.")
    m.figura(IMG / "fig_ex_harmonico.png", "8.2",
             f"Seção EBG para supressão de 2f₀: rejeição de cerca de {n(-h['s21_2f0'], 0)} dB em 4,9 GHz, perda pequena em "
             "2,45 GHz e nenhuma banda em 3f₀ (perfil senoidal).", 12)
    m.p(f"**Discussão.** A reflexão de {n(h['s11_f0'], 1)} dB em f₀ vem das ondulações da banda passante, típicas de "
        "estruturas periódicas truncadas. Num circuito ativo, essa reflexão altera a carga vista pelo "
        "transistor e precisa entrar no projeto do casamento. O livro sugere *tapering* das pontas "
        "(Hamming) para reduzir as ondulações. Se também fosse preciso atenuar 3f₀, um perfil com duas "
        "senoides somadas, sintonizadas em 2f₀ e 3f₀, resolveria (pp. 88–91 e 97).")
    # ---- miniaturizacao
    m.h("8.3  Miniaturização de uma linha de 90° a 1 GHz", 2)
    m.p("**Especificação.** Uma rede de casamento usa um trecho λ/4 de 50 Ω a 1 GHz em RO4003C. Deseja-se "
        f"reduzir seu comprimento com uma linha *slow-wave* de {mi['n']} células, mantendo 50 Ω e afastando "
        f"a primeira espúria para f_s = {n(mi['fs'] / 1e9, 0)} GHz.")
    m.p(f"**Linha convencional.** W = {n(mi['w50'], 2)} mm, εeff = {n(mi['ee50'], 2)} e λg/4 = "
        f"**{n(mi['comp_conv'], 1)} mm**.")
    m.p(f"**Projeto slow-wave.** Cada célula deve ter φ = 90°/{mi['n']} = {n(mi['fase'], 1)}° a 1 GHz. Pelas "
        f"Eqs. (5.5) e (5.4): L = {n(mi['L'], 3)} nH, C = {n(mi['C'], 3)} pF e C_ls = {n(mi['Cls'], 3)} pF. "
        f"A linha descarregada tem Z₀ = {n(mi['zo'], 1)} Ω, f_c = {n(mi['fc'], 2)} GHz e o fator slow-wave "
        f"é {n(mi['sw'], 2)}. Na geometria (Z_p = {mi['zp']} Ω): período {n(mi['periodo'], 2)} mm, linha de "
        f"{n(mi['w_linha'], 2)} mm e patch de {n(mi['w_patch'], 2)} × {n(mi['l_patch'], 2)} mm.")
    m.tabela(["Grandeza", "Convencional", "Slow-wave (4 células)"],
             [["Comprimento", f"{n(mi['comp_conv'], 1)} mm", f"{n(mi['comp_sw'], 1)} mm"],
              ["Fase a 1 GHz", "90°", f"{n(mi['fase_total'], 1)}° (modelo distribuído)"],
              ["|S₁₁| a 1 GHz (portas de 50 Ω)", "–", f"{n(mi['s11'], 1)} dB"],
              ["Banda espúria", "2 GHz (λ/2)", f"corte em {n(mi['fc'], 1)} GHz; espúria só perto de 12 GHz"]],
             [5.4, 5.0, 5.6], "8.1", "Comparação entre a linha λ/4 convencional e a versão slow-wave.")
    m.p(f"**Discussão.** A redução de comprimento é de cerca de **{n(mi['reducao'], 0)}%**, em linha com os "
        "30% do exemplo CPW do livro. O ganho vem acompanhado de largura: os patches ocupam cerca de "
        f"{n(mi['w_patch'], 1)} mm de largura. A linha também passou a ser um passa-baixas, o que pode ser "
        "uma vantagem (filtragem de harmônicos de graça) ou uma restrição (banda limitada a f_c). Ganhos "
        "maiores exigem carregar mais (fator slow-wave maior), o que pede linhas descarregadas de impedância "
        "mais alta e, portanto, mais finas.")


# =============================================================================
def cap9(m):
    m.h("9  Limitações, validação e próximos passos")
    m.h("9.1  Validação do software", 2)
    m.p("O LTA possui testes automatizados (`python -m unittest discover testes`) que comparam os cálculos "
        "com valores publicados no livro:")
    m.tabela(["Item", "Livro", "LTA"],
             [["W de 50 Ω em RO3010 (h = 1,27 mm)", "1,2 mm", "1,186 mm"],
              ["Z₀ máx/mín, perfil senoidal |K₁| = 0,64 cm⁻¹", "75 / 33,3 Ω", "75,1 / 33,3 Ω"],
              ["|S₂₁|min senoidal (Eq. 6.5)", "−49,6 dB", "−49,6 dB"],
              ["Banda rejeitada senoidal (Eq. 6.6)", "2,55 GHz", "2,61 GHz"],
              ["|S₂₁|min quadrada, n = 1, 3, 5", "−49,3 / −12,5 / −5,7 dB", "−49,3 / −12,5 / −5,7 dB"],
              ["f_max (Eq. 6.5)", "5,66 GHz", "5,74 GHz"],
              ["Exemplo da Fig. 2.34: L, C, C_ls", "3,1 nH; 0,17 pF; 1,2 pF", "3,30 nH; 0,157 pF; 1,31 pF"]],
             [7.0, 4.5, 4.5], "9.1", "Comparação com valores publicados em Martín (2015).")
    m.h("9.2  Limitações conhecidas", 2)
    m.lista(["Microstrip quase-estática (Hammerstad-Jensen): espessura de cobre nula, sem dispersão e sem "
             "perdas nos cálculos de S.",
             "Modos acoplados e cascata ABCD desprezam descontinuidades (degraus), radiação e acoplamento "
             "a modos de ordem superior.",
             "A geometria da linha com carga capacitiva é estimativa inicial (Eq. 5.7), a ser otimizada em "
             "simulador eletromagnético.",
             "EBG com furos no terra: apenas geometria e estimativa de Bragg.",
             "Os arquivos para HFSS, CST e KiCad ainda precisam ser testados nas ferramentas instaladas."])
    m.h("9.3  Próximos passos sugeridos", 2)
    m.p("O capítulo 2 do livro prepara o terreno para os capítulos seguintes, que tratam das linhas "
        "metamateriais. A aba Célula ABCD já calcula a dispersão de células CRLH (Ls, Cs, Lp, Cp), base dos "
        "componentes de dupla banda e das antenas *leaky-wave* (" + LIVRO + ", seções 3.4 e 4.2, pp. 149–"
        "162 e 215–290). Extensões naturais do software são: projeto de células CRLH balanceadas, síntese "
        "inversa de perfis EBG (da resposta desejada para K(z)) e inclusão de perdas de condutor e "
        "dielétrico.")


def referencias(m):
    m.h("Referências")
    m.lista(["MARTÍN, F. **Artificial Transmission Lines for RF and Microwave Applications**. Hoboken: "
             "John Wiley & Sons, 2015. (Wiley Series in Microwave and Optical Engineering). "
             "ISBN 978-1-118-48760-0. Capítulos 1 e 2.",
             "HAMMERSTAD, E.; JENSEN, O. Accurate Models for Microstrip Computer-Aided Design. In: **IEEE "
             "MTT-S International Microwave Symposium Digest**, 1980, p. 407–409.",
             "POZAR, D. M. **Microwave Engineering**. 4. ed. Hoboken: John Wiley & Sons, 2012. Seção 4.4 "
             "(matriz ABCD e conversão para parâmetros S)."])
    m.p("As referências originais dos exemplos de aplicação (refletores, ressonadores de Bragg, filtros "
        "*wiggly-line*, antena ativa e linhas de atraso chirpadas) estão listadas ao final do capítulo 2 do "
        "livro, pp. 114–118.")


def main():
    print("Gerando figuras e números...")
    F = figuras.gerar_tudo()
    faltando = [p.name for p in IMG.glob("tela_*.png")]
    if len(faltando) < 10:
        print("Aviso: capturas de tela ausentes; rode doc/fontes/capturar_telas.py antes.")
    m = Manual("LTA – Linhas de Transmissão Artificiais Periódicas | Material didático")
    capa(m)
    m.sumario()
    cap1(m)
    cap2(m, F)
    cap3(m, F)
    cap4(m, F)
    cap5(m, F)
    cap6(m, F)
    cap7(m)
    cap8(m, F)
    cap9(m)
    referencias(m)
    m.cabecalho_rodape()
    m.salvar(SAIDA)
    print("Documento:", SAIDA.relative_to(RAIZ))
    atualizar_com_word(SAIDA)


def atualizar_com_word(caminho):
    """Atualiza o sumário pelo Word e gera um PDF ao lado (opcional, requer pywin32)."""
    try:
        import win32com.client as wc
    except ImportError:
        print("pywin32 ausente: abra o .docx e atualize o sumário manualmente (F9).")
        return
    w = wc.dynamic.Dispatch("Word.Application")
    w.Visible = False
    try:
        d = w.Documents.Open(str(caminho))
        d.TablesOfContents(1).Update()
        d.Save()
        d.SaveAs2(str(caminho.with_suffix(".pdf")), 17)
        d.Close(0)
        print("Sumário atualizado e PDF gerado.")
    finally:
        w.Quit()


if __name__ == "__main__":
    main()
