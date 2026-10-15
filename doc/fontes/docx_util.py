"""Funções auxiliares para montar o manual em Word (python-docx + equações OMML nativas)."""
import re
from pathlib import Path

import latex2mathml.converter
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

# Conversor MathML -> OMML que acompanha o Microsoft Office
_XSL_CANDIDATOS = [
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
    Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSL = next((p for p in _XSL_CANDIDATOS if p.exists()), None)
if _XSL is None:
    raise SystemExit("MML2OMML.XSL não encontrado: é necessário o Microsoft Office instalado.")
_TRANSF = etree.XSLT(etree.parse(str(_XSL)))

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
CINZA = RGBColor(0x59, 0x59, 0x59)
LARGURA_TEXTO_CM = 16.0


_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_BMATRIX = re.compile(r"\\begin\{bmatrix\}(.*?)\\end\{bmatrix\}", re.S)


def _m(tag, **attrs):
    el = etree.Element(f"{{{_M}}}{tag}")
    for k, v in attrs.items():
        el.set(f"{{{_M}}}{k}", v)
    return el


def _matriz(conteudo):
    """bmatrix -> objeto matriz nativo do Word entre colchetes extensíveis."""
    d = _m("d")
    dpr = _m("dPr")
    dpr.append(_m("begChr", val="["))
    dpr.append(_m("endChr", val="]"))
    d.append(dpr)
    e = _m("e")
    mat = _m("m")
    for linha in conteudo.split(r"\\"):
        mr = _m("mr")
        for celula in linha.split("&"):
            ce = _m("e")
            for filho in _omml_simples(celula.strip()):
                ce.append(filho)
            mr.append(ce)
        mat.append(mr)
    e.append(mat)
    d.append(e)
    return d


def _omml_simples(latex):
    mml = latex2mathml.converter.convert(latex)
    return _TRANSF(etree.fromstring(mml)).getroot()


def omml(latex):
    # o conversor descarta \quad/\qquad; espaços em \text{} são preservados
    latex = latex.replace(r"\qquad", r"\text{      }").replace(r"\quad", r"\text{   }")
    # matrizes: substitui por marcadores e insere o objeto nativo depois
    matrizes = []

    def marcar(mt):
        matrizes.append(mt.group(1))
        return r"\mathrm{MATRIZ" + "ABCDEFGHIJ"[len(matrizes) - 1] + "}"

    latex = _BMATRIX.sub(marcar, latex)
    raiz = _omml_simples(latex)
    for i, conteudo in enumerate(matrizes):
        marca = "MATRIZ" + "ABCDEFGHIJ"[i]
        for r in raiz.iter(f"{{{_M}}}r"):
            t = r.find(f"{{{_M}}}t")
            if t is not None and t.text == marca:
                r.getparent().replace(r, _matriz(conteudo))
                break
    return raiz


# _{...} ou _palavra -> subscrito; ^{...} ou ^(...) -> sobrescrito
_INDICES = re.compile(r"_\{([^}]+)\}|_([A-Za-z0-9áéíóúâêôãõç]+)|\^\{([^}]+)\}|\^\(([^)]+)\)")


class Manual:
    def __init__(self, titulo_cabecalho):
        self.doc = Document()
        self._estilos()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
        self.titulo_cabecalho = titulo_cabecalho

    # ------------------------------------------------------------------ estilos
    def _estilos(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for nome, tam, antes in (("Heading 1", 17, 0), ("Heading 2", 13.5, 14), ("Heading 3", 11.5, 10)):
            s = st[nome]
            s.font.name = "Calibri"
            s.font.size = Pt(tam)
            s.font.bold = True
            s.font.color.rgb = AZUL
            s.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
            s.paragraph_format.space_before = Pt(antes)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        st["Heading 1"].paragraph_format.page_break_before = True
        cap = st["Caption"]
        cap.font.size = Pt(9.5)
        cap.font.italic = True
        cap.font.bold = False
        cap.font.color.rgb = CINZA
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)

    # ---------------------------------------------------------------- cabeçalho
    def cabecalho_rodape(self):
        for sec in self.doc.sections:
            sec.different_first_page_header_footer = True
            hp = sec.header.paragraphs[0]
            hp.text = self.titulo_cabecalho
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in hp.runs:
                r.font.size = Pt(8.5)
                r.font.color.rgb = CINZA
            fp = sec.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = fp.add_run()
            r.font.size = Pt(9)
            self._campo(r, "PAGE")

    @staticmethod
    def _campo(run, instrucao, texto="1"):
        for tipo, conteudo in (("begin", None), ("instr", instrucao), ("separate", None), ("texto", texto),
                               ("end", None)):
            if tipo == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {conteudo} "
            elif tipo == "texto":
                el = OxmlElement("w:t")
                el.text = conteudo
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tipo)
            run._r.append(el)

    def sumario(self):
        self.h("Sumário", 1, numerar=False)
        p = self.doc.add_paragraph()
        self._campo(p.add_run(), 'TOC \\o "1-2" \\h \\z \\u',
                    "Clique com o botão direito e escolha 'Atualizar campo' para gerar o sumário.")

    # ------------------------------------------------------------------ texto
    def h(self, texto, nivel=1, numerar=True):
        return self.doc.add_heading(texto, level=nivel)

    @staticmethod
    def _com_indices(par, texto, negrito=False, italico=False, tamanho=None):
        """Escreve texto convertendo X_sub e X^{sup} em subscrito/sobrescrito reais."""
        pos = 0
        pedacos = []
        for mt in _INDICES.finditer(texto):
            pedacos.append((texto[pos:mt.start()], None))
            sub = mt.group(1) or mt.group(2)
            pedacos.append((sub, "sub") if sub else (mt.group(3) or mt.group(4), "sup"))
            pos = mt.end()
        pedacos.append((texto[pos:], None))
        for txt, tipo in pedacos:
            if not txt:
                continue
            r = par.add_run(txt)
            r.bold = negrito or None
            r.italic = italico or None
            if tamanho:
                r.font.size = Pt(tamanho)
            if tipo == "sub":
                r.font.subscript = True
            elif tipo == "sup":
                r.font.superscript = True

    def _runs(self, par, texto, negrito=False, tamanho=None):
        """Marcação simples: **negrito**, *itálico*, `código`, X_sub, X^{sup}."""
        partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto)
        for parte in partes:
            if not parte:
                continue
            if parte.startswith("**"):
                self._com_indices(par, parte[2:-2], True, False, tamanho)
            elif parte.startswith("*"):
                self._com_indices(par, parte[1:-1], negrito, True, tamanho)
            elif parte.startswith("`"):
                r = par.add_run(parte[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(tamanho - 1 if tamanho else 9.5)
            else:
                self._com_indices(par, parte, negrito, False, tamanho)

    def p(self, texto, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY):
        par = self.doc.add_paragraph()
        par.alignment = alinhamento
        self._runs(par, texto)
        return par

    def lista(self, itens, numerada=False):
        estilo = "List Number" if numerada else "List Bullet"
        for it in itens:
            par = self.doc.add_paragraph(style=estilo)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.space_after = Pt(3)
            self._runs(par, it)

    def eq(self, latex, rotulo):
        """Equação em modo de exibição (nativa do Word) com número à direita.

        Usa uma tabela sem bordas de duas colunas: com o número no mesmo
        parágrafo, o Word trataria a equação como "em linha" (frações reduzidas).
        """
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_eq, c_num = t.rows[0].cells
        larguras = (LARGURA_TEXTO_CM - 1.6, 1.6)
        for c, w in zip((c_eq, c_num), larguras):
            c.width = Cm(w)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par = c_eq.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(3)
        par.paragraph_format.space_after = Pt(3)
        bloco = OxmlElement("m:oMathPara")
        bloco.append(omml(latex))
        par._p.append(bloco)
        pn = c_num.paragraphs[0]
        pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pn.paragraph_format.space_after = Pt(0)
        pn.add_run(f"({rotulo})")
        respiro = self.doc.add_paragraph()
        respiro.paragraph_format.space_after = Pt(0)
        respiro.paragraph_format.line_spacing = Pt(5)
        return t

    def eq_inline(self, partes):
        """Parágrafo misturando texto e equações: partes = ['texto', ('$', latex), ...]."""
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for parte in partes:
            if isinstance(parte, tuple):
                par._p.append(omml(parte[1]))
            else:
                self._runs(par, parte)
        return par

    def nota(self, texto, titulo="Observação"):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ppr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        esq = OxmlElement("w:left")
        for k, v in (("val", "single"), ("sz", "24"), ("space", "8"), ("color", "2E75B6")):
            esq.set(qn(f"w:{k}"), v)
        bdr.append(esq)
        ppr.append(bdr)
        shd = OxmlElement("w:shd")
        for k, v in (("val", "clear"), ("color", "auto"), ("fill", "EAF1FB")):
            shd.set(qn(f"w:{k}"), v)
        ppr.append(shd)
        par.paragraph_format.left_indent = Cm(0.4)
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(10)
        r = par.add_run(f"{titulo}. ")
        r.bold = True
        r.font.color.rgb = AZUL
        self._runs(par, texto)
        return par

    def figura(self, caminho, rotulo, legenda, largura_cm=LARGURA_TEXTO_CM):
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(str(caminho), width=Cm(largura_cm))
        cap = self.doc.add_paragraph(style="Caption")
        r = cap.add_run(f"Figura {rotulo} – ")
        r.bold = True
        self._runs(cap, legenda)

    def tabela(self, cabecalho, linhas, larguras_cm, rotulo=None, legenda=None, fonte=9.5):
        if rotulo:
            cap = self.doc.add_paragraph(style="Caption")
            cap.paragraph_format.keep_with_next = True
            r = cap.add_run(f"Tabela {rotulo} – ")
            r.bold = True
            self._runs(cap, legenda)
        t = self.doc.add_table(rows=1, cols=len(cabecalho))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, txt in enumerate(cabecalho):
            c = t.rows[0].cells[i]
            c.text = ""
            par = c.paragraphs[0]
            self._runs(par, txt, negrito=True, tamanho=fonte)
            shd = OxmlElement("w:shd")
            for k, v in (("val", "clear"), ("color", "auto"), ("fill", "D9E2F3")):
                shd.set(qn(f"w:{k}"), v)
            c._tc.get_or_add_tcPr().append(shd)
        for linha in linhas:
            cells = t.add_row().cells
            for i, txt in enumerate(linha):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                if isinstance(txt, tuple):
                    par._p.append(omml(txt[1]))
                else:
                    self._runs(par, str(txt), tamanho=fonte)
        for row in t.rows:
            for i, w in enumerate(larguras_cm):
                row.cells[i].width = Cm(w)
        # repetir cabeçalho em quebra de página
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def quebra(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def salvar(self, caminho):
        self.doc.save(str(caminho))
